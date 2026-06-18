"""Native-file builders.

Each builder turns an :class:`~generator.model.Attachment` spec into real,
openable file bytes *and* the plain text a processing engine would extract
from it. We take care to make the bytes deterministic (fixed timestamps, no
random ids) so that re-running the generator does not churn every hash.
"""
from __future__ import annotations

import datetime as dt
import io
import os
import re
import zipfile

# Make reportlab output byte-stable (fixed date + document id) for reproducible
# hashes. This env var must be set before reportlab is imported.
os.environ.setdefault("RL_invariant", "1")

from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from . import cast
from .model import Attachment

# A constant zip timestamp (1980-01-01) for reproducible Office files.
_ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)
_MODIFIED_RE = re.compile(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)")


def _normalise_zip(data: bytes, modified: dt.datetime | None = None) -> bytes:
    """Rewrite a zip (docx/xlsx) deterministically.

    Members are written in sorted order with a constant timestamp. If
    ``modified`` is given, the document's ``dcterms:modified`` property is
    forced to that value — necessary because openpyxl otherwise stamps the
    current wall-clock time on save, which would change the bytes (and hash)
    on every build.
    """
    stamp = (modified.strftime("%Y-%m-%dT%H:%M:%SZ").encode() if modified else None)
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name in sorted(src.namelist()):
            content = src.read(name)
            if stamp and name == "docProps/core.xml":
                content = _MODIFIED_RE.sub(rb"\g<1>" + stamp + rb"\g<2>", content)
            info = zipfile.ZipInfo(name, date_time=_ZIP_EPOCH)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            zf.writestr(info, content)
    return out.getvalue()


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for path in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ):
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default()


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------
def _build_docx(att: Attachment) -> tuple[bytes, str]:
    author = cast.get(att.author_key).name
    doc = Document()
    text_parts: list[str] = []

    if att.confidentiality:
        banner = doc.add_paragraph()
        run = banner.add_run(att.confidentiality.upper())
        run.bold = True
        run.font.color.rgb = RGBColor(0x99, 0x00, 0x00)
        run.font.size = Pt(9)
        text_parts.append(att.confidentiality.upper())

    heading = doc.add_heading(att.title, level=1)
    heading.alignment = 1
    text_parts.append(att.title)

    for para in att.spec.get("paragraphs", []):
        doc.add_paragraph(para)
        text_parts.append(para)

    table_spec = att.spec.get("table")
    if table_spec:
        headers = table_spec["headers"]
        rows = table_spec["rows"]
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Grid Accent 1"
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = str(h)
        text_parts.append("\t".join(str(h) for h in headers))
        for r in rows:
            cells = table.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = str(v)
            text_parts.append("\t".join(str(v) for v in r))

    closing = att.spec.get("closing")
    if closing:
        doc.add_paragraph()
        doc.add_paragraph(closing)
        text_parts.append(closing)

    cp = doc.core_properties
    cp.author = author
    cp.last_modified_by = att.spec.get("last_modified_by", author)
    cp.created = att.created
    cp.modified = att.modified
    cp.title = att.title
    cp.category = att.storyline

    buf = io.BytesIO()
    doc.save(buf)
    return _normalise_zip(buf.getvalue(), att.modified), "\n".join(text_parts)


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------
def _build_xlsx(att: Attachment) -> tuple[bytes, str]:
    author = cast.get(att.author_key).name
    wb = Workbook()
    ws = wb.active
    ws.title = att.spec.get("sheet", "Sheet1")[:31]
    text_parts: list[str] = []

    headers = att.spec.get("headers", [])
    rows = att.spec.get("rows", [])

    title_cell = ws.cell(row=1, column=1, value=att.title)
    title_cell.font = Font(bold=True, size=13)
    text_parts.append(att.title)
    start = 3

    if headers:
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=start, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center")
        text_parts.append("\t".join(str(h) for h in headers))
        start += 1

    for r, row in enumerate(rows, start=start):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
        text_parts.append("\t".join(str(v) for v in row))

    # Sensible column widths.
    for c in range(1, max(1, len(headers)) + 1):
        ws.column_dimensions[chr(64 + c)].width = 22

    note = att.spec.get("note")
    if note:
        ws.cell(row=start + len(rows) + 1, column=1, value=note)
        text_parts.append(note)

    props = wb.properties
    props.creator = author
    props.lastModifiedBy = att.spec.get("last_modified_by", author)
    props.created = att.created
    props.modified = att.modified
    props.title = att.title

    buf = io.BytesIO()
    wb.save(buf)
    return _normalise_zip(buf.getvalue(), att.modified), "\n".join(text_parts)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def _build_pdf(att: Attachment) -> tuple[bytes, str]:
    author = cast.get(att.author_key).name
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=LETTER,
        title=att.title, author=author, subject=att.storyline,
        creator="Royal Scriptorium",
        leftMargin=1 * inch, rightMargin=1 * inch,
        topMargin=1 * inch, bottomMargin=1 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Title"], fontName="Times-Bold", fontSize=16)
    body_style = ParagraphStyle("body", parent=styles["Normal"], fontName="Times-Roman", fontSize=11, leading=16, spaceAfter=10)
    conf_style = ParagraphStyle("conf", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8, textColor="#990000", alignment=TA_CENTER, spaceAfter=14)

    flow: list = []
    text_parts: list[str] = []
    if att.confidentiality:
        flow.append(Paragraph(att.confidentiality.upper(), conf_style))
        text_parts.append(att.confidentiality.upper())
    flow.append(Paragraph(att.title, title_style))
    flow.append(Spacer(1, 18))
    text_parts.append(att.title)

    for para in att.spec.get("paragraphs", []):
        flow.append(Paragraph(para, body_style))
        text_parts.append(para)

    table_spec = att.spec.get("table")
    if table_spec:
        data = [table_spec["headers"], *[[str(c) for c in r] for r in table_spec["rows"]]]
        tbl = Table(data, hAlign="LEFT")
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), "#4472C4"),
            ("TEXTCOLOR", (0, 0), (-1, 0), "#FFFFFF"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, "#999999"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), ["#FFFFFF", "#F2F2F2"]),
        ]))
        flow.append(Spacer(1, 10))
        flow.append(tbl)
        text_parts.append("\t".join(str(h) for h in table_spec["headers"]))
        for r in table_spec["rows"]:
            text_parts.append("\t".join(str(c) for c in r))

    closing = att.spec.get("closing")
    if closing:
        flow.append(Spacer(1, 18))
        flow.append(Paragraph(closing, body_style))
        text_parts.append(closing)

    doc.build(flow)
    return buf.getvalue(), "\n".join(text_parts)


# --------------------------------------------------------------------------
# PNG  (stylised placeholders — never real imagery)
# --------------------------------------------------------------------------
def _build_png(att: Attachment) -> tuple[bytes, str]:
    style = att.spec.get("style", "parchment")
    caption = att.spec.get("caption", att.title)
    subtitle = att.spec.get("subtitle", "")
    W, H = 760, 980

    if style == "portrait":
        img = Image.new("RGB", (W, H), (38, 28, 18))
        d = ImageDraw.Draw(img)
        d.rectangle([28, 28, W - 28, H - 28], outline=(196, 160, 64), width=10)
        d.rectangle([70, 70, W - 70, H - 70], fill=(70, 84, 96))
        # A simple silhouette stands in for a portrait — no real likeness.
        d.ellipse([W // 2 - 90, 230, W // 2 + 90, 410], fill=(26, 32, 38))
        d.polygon([(W // 2 - 150, 640), (W // 2 + 150, 640), (W // 2 + 110, 430),
                   (W // 2 - 110, 430)], fill=(26, 32, 38))
        plate_y = H - 150
    elif style == "seal":
        img = Image.new("RGB", (W, H), (244, 238, 222))
        d = ImageDraw.Draw(img)
        cx, cy = W // 2, H // 2
        d.ellipse([cx - 220, cy - 220, cx + 220, cy + 220], fill=(140, 24, 24), outline=(90, 12, 12), width=8)
        d.ellipse([cx - 150, cy - 150, cx + 150, cy + 150], outline=(220, 200, 160), width=6)
        plate_y = cy + 260
    else:  # parchment / scanned document look
        img = Image.new("RGB", (W, H), (243, 233, 208))
        d = ImageDraw.Draw(img)
        for i in range(0, H, 4):
            shade = 233 - (i % 24)
            d.line([(0, i), (W, i)], fill=(shade, shade - 8, shade - 28))
        d.rectangle([40, 40, W - 40, H - 40], outline=(120, 96, 60), width=3)
        plate_y = 120

    font_title = _font(34)
    font_sub = _font(20)
    text = caption if style != "parchment" else att.title
    tw = d.textlength(text, font=font_title)
    d.text(((W - tw) / 2, plate_y), text, fill=(20, 18, 14), font=font_title)
    if subtitle:
        sw = d.textlength(subtitle, font=font_sub)
        d.text(((W - sw) / 2, plate_y + 48), subtitle, fill=(60, 50, 36), font=font_sub)

    # Parchment style also renders body lines so it has extractable text.
    extract = f"[Image] {caption}".strip()
    if style == "parchment":
        y = 200
        body_font = _font(18)
        lines = att.spec.get("lines", [])
        for ln in lines:
            d.text((80, y), ln, fill=(30, 24, 16), font=body_font)
            y += 34
        if lines:
            extract = "\n".join([att.title, *lines])

    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=False)
    return buf.getvalue(), extract


# --------------------------------------------------------------------------
# TXT
# --------------------------------------------------------------------------
def _build_txt(att: Attachment) -> tuple[bytes, str]:
    text = att.spec.get("text", "")
    return text.encode("utf-8"), text


def _build_binary(att: Attachment) -> tuple[bytes, str]:
    """Pass-through for bytes generated elsewhere (audio, RSMF chat containers)."""
    return att.spec["data"], att.spec.get("text", "")


_BUILDERS = {
    "docx": _build_docx,
    "xlsx": _build_xlsx,
    "pdf": _build_pdf,
    "png": _build_png,
    "txt": _build_txt,
    "binary": _build_binary,
}


def build(att: Attachment) -> None:
    """Populate ``att.data`` and ``att.text`` in place."""
    builder = _BUILDERS.get(att.kind)
    if builder is None:
        raise ValueError(f"unknown attachment kind: {att.kind!r}")
    att.data, att.text = builder(att)
